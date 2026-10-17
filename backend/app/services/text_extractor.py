"""Secure Text Extraction Service for PDF, DOCX, and Notes."""

import re
from io import BytesIO
from typing import Any, Dict, Tuple
from pypdf import PdfReader
from docx import Document as DocxDocument

from ..logger import logger
from ..schemas.document import FileType
from .document_validator import MalformedFileError, sanitize_and_validate_notes


def clean_extracted_text(text: str) -> str:
    """Sanitizes extracted text and normalizes spacing."""
    if not text:
        return ""
    # Remove null bytes
    cleaned = text.replace("\x00", "")
    # Normalize multiple blank lines to double newline
    cleaned = re.sub(r"\n{3,}", "\n\n", cleaned)
    # Strip leading/trailing whitespace
    return cleaned.strip()


def extract_pdf_text(content: bytes) -> Tuple[str, Dict[str, Any]]:
    """Extracts text and metadata from PDF bytes."""
    try:
        reader = PdfReader(BytesIO(content))
        if reader.is_encrypted:
            # Try to decrypt with empty password (common for readable encrypted PDFs)
            try:
                reader.decrypt("")
            except Exception:
                raise MalformedFileError("PDF is encrypted and password-protected")

        page_count = len(reader.pages)
        page_texts = []

        for idx, page in enumerate(reader.pages):
            try:
                text = page.extract_text() or ""
                page_texts.append(text)
            except Exception as e:
                logger.warning(f"Error extracting page {idx + 1} text: {e}")
                page_texts.append("")

        full_text = clean_extracted_text("\n\n".join(page_texts))
        word_count = len(re.findall(r"\b\w+\b", full_text))

        metadata: Dict[str, Any] = {
            "page_count": page_count,
            "word_count": word_count,
            "character_count": len(full_text),
            "is_encrypted": reader.is_encrypted,
        }

        # Extract document info if present
        if reader.metadata:
            meta = reader.metadata
            if meta.title:
                metadata["pdf_title"] = meta.title[:200]
            if meta.author:
                metadata["pdf_author"] = meta.author[:100]

        return full_text, metadata

    except MalformedFileError:
        raise
    except Exception as e:
        logger.error(f"Failed to extract PDF text: {str(e)}")
        raise MalformedFileError(f"Failed to parse PDF document structure: {str(e)}")


def extract_docx_text(content: bytes) -> Tuple[str, Dict[str, Any]]:
    """Extracts text and metadata from DOCX bytes."""
    try:
        doc = DocxDocument(BytesIO(content))
        paragraph_texts = [p.text for p in doc.paragraphs if p.text.strip()]

        table_texts = []
        for table in doc.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_cells:
                    table_texts.append(" | ".join(row_cells))

        combined_text = "\n\n".join(paragraph_texts)
        if table_texts:
            combined_text += "\n\n" + "\n".join(table_texts)

        full_text = clean_extracted_text(combined_text)
        word_count = len(re.findall(r"\b\w+\b", full_text))

        metadata: Dict[str, Any] = {
            "paragraph_count": len(doc.paragraphs),
            "table_count": len(doc.tables),
            "word_count": word_count,
            "character_count": len(full_text),
        }

        return full_text, metadata

    except Exception as e:
        logger.error(f"Failed to extract DOCX text: {str(e)}")
        raise MalformedFileError(f"Failed to parse DOCX document structure: {str(e)}")


def extract_notes_text(content: str) -> Tuple[str, Dict[str, Any]]:
    """Extracts and sanitizes raw notes text."""
    sanitized = sanitize_and_validate_notes(content)
    cleaned = clean_extracted_text(sanitized)
    word_count = len(re.findall(r"\b\w+\b", cleaned))
    lines = cleaned.splitlines()

    metadata: Dict[str, Any] = {
        "line_count": len(lines),
        "word_count": word_count,
        "character_count": len(cleaned),
    }

    return cleaned, metadata


def extract_text(file_type: FileType, content: Any) -> Tuple[str, Dict[str, Any]]:
    """Unified dispatcher for extracting text according to FileType."""
    if file_type == FileType.PDF:
        if not isinstance(content, bytes):
            raise ValueError("Expected bytes for PDF extraction")
        return extract_pdf_text(content)
    elif file_type == FileType.DOCX:
        if not isinstance(content, bytes):
            raise ValueError("Expected bytes for DOCX extraction")
        return extract_docx_text(content)
    elif file_type == FileType.NOTES:
        if not isinstance(content, str):
            raise ValueError("Expected str for notes extraction")
        return extract_notes_text(content)
    else:
        raise ValueError(f"Unsupported file type for extraction: {file_type}")
